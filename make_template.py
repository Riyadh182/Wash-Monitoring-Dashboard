"""Builds the Google-Sheet template (.xlsx) for one unit: python make_template.py CWL-1 ../Wash_Pilot_Template_CWL-1.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL
from common import *

unit = sys.argv[1] if len(sys.argv) > 1 else 'CWL-1'
out = sys.argv[2] if len(sys.argv) > 2 else f'Wash_Pilot_Template_{unit}.xlsx'
if not CFG['units'][unit]['dry_verified'] and CFG['units'][unit]['dry_processes']:
    print(f'WARNING: dry-process capacity for {unit} duplicates an earlier unit in Capacity_August_2026 - verify before use.')
if not CFG['units'][unit]['dry_processes']:
    print(f'NOTE: no dry-process capacity in source for {unit}; Master_Process will be empty - fill it in.')

F = lambda b=False, c='222222', s=10: Font(name='Arial', size=s, bold=b, color=c)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
NAVY, TEAL, KEY, INP = '1B2A41', '2A9D8F', 'ECEFF3', 'FFF8CC'
wb = Workbook()

def head(ws, cols, key_n=0, widths=None, row=1):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row, j, h); c.font = F(True, 'FFFFFF', 10); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.fill = fill(NAVY if j <= key_n else TEAL)
        ws.column_dimensions[CL(j)].width = (widths or {}).get(h, 16)
    ws.row_dimensions[row].height = 32; ws.freeze_panes = ws.cell(row + 1, 1)

# ---- README
r = wb.active; r.title = 'README'; r.sheet_view.showGridLines = False
r.column_dimensions['A'].width = 3; r.column_dimensions['B'].width = 120
lines = [('WASH & DRY HOURLY LOG - ' + unit, True), ('', False),
 ('1. Upload this file to Google Drive and open with Google Sheets (File > Save as Google Sheets).', False),
 ('2. Extensions > Apps Script > paste apps_script/Code.gs > Save > reload. A "Wash Tools" menu appears.', False),
 ('3. Each production day: Wash Tools > Create rows for a day. It fills Date, Hour, Unit and Machine/Process for all 24 hours.', False),
 ('4. Supervisors type only numbers (Input, Output, QC, Defect, Rewash, Downtime). Grey columns are keys - do not edit.', False),
 ('5. Style: pick from the dropdown (list in the Lists tab) or type a new one. Wash type and Downtime reason are dropdowns.', False),
 ('6. WIP_Log: one row per stage per day with the physical opening WIP. Stage = WASH or a dry process name (exactly as in Master_Process).', False),
 ('7. Master_Machine / Master_Process: plan per hour comes from here. Edit Plan_per_hr for a machine, or set Active = N for a machine that is down for the day.', False),
 ('8. Share this Google Sheet with the dashboard service account as VIEWER. Floor staff get EDITOR access.', False),
 ('9. Month end: File > Make a copy (name it with the month), keep it as the archive, then clear Wash_Log/Dry_Log/WIP_Log data rows (keep headers) for the new month.', False),
 ('', False), ('Rows with "EXAMPLE" in Remarks are ignored by the dashboard - delete them whenever you like.', False),
 ('Hour slots assume the production day starts at Day_Start_Hour in Settings (08:00 default): hours after midnight belong to the same production date.', False)]
for i, (t, b) in enumerate(lines, 2):
    c = r.cell(i, 2, t); c.font = F(b, NAVY if b else '222222', 14 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')

# ---- Settings
s = wb.create_sheet('Settings'); head(s, ['Key', 'Value', 'Meaning'], widths={'Key': 20, 'Value': 12, 'Meaning': 90})
notes = {'Run_Hours': 'Hours the plant runs per day (hourly plan = capacity per day / this)', 'Day_Start_Hour': 'Production day start hour (0-23); shifts A/B/C are 8 hours each',
         'Plan_Basis': '1 = plan at 100% of capacity; 0.85 = plan at 85% load', 'Target_Pct': 'Achievement at/above = green', 'Alert_Pct': 'Achievement below = red',
         'Max_Defect_Pct': 'Defect limit (defect pcs / QC inspected pcs) - placeholder, set yours', 'Max_Rewash_Pct': 'Rewash limit (rewash pcs / output pcs) - placeholder, set yours'}
for i, (k, v) in enumerate(DEFAULT_SETTINGS, 2):
    s.cell(i, 1, k).font = F(True); c = s.cell(i, 2, v); c.font = F(c='0000FF'); c.fill = fill(INP); s.cell(i, 3, notes[k]).font = F(s=9, c='666666')
    if k.endswith('Pct') or k == 'Plan_Basis': c.number_format = '0%'

# ---- Masters
mm = wb.create_sheet('Master_Machine'); head(mm, ['Unit', 'Machine_ID', 'Type', 'Plan_per_hr', 'Active'], key_n=3)
for i, m in enumerate(machines_for(unit), 2):
    for j, k in enumerate(['Unit', 'Machine_ID', 'Type', 'Plan_per_hr', 'Active'], 1):
        c = mm.cell(i, j, m[k]); c.font = F(c='0000FF' if j >= 4 else '222222')
        if j >= 4: c.fill = fill(INP)
dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=False); mm.add_data_validation(dv); dv.add(f'E2:E{mm.max_row}')
mm['G1'] = 'Plan_per_hr = unit wash capacity per day / machines / run hours (Capacity_August_2026). It is an average - overwrite per machine when you know the real capacity.'
mm['G1'].font = F(s=9, c='666666'); mm.column_dimensions['G'].width = 80

mp = wb.create_sheet('Master_Process'); head(mp, ['Unit', 'Process', 'Machines', 'Capacity_per_day', 'Active'], key_n=2, widths={'Process': 28})
for i, p in enumerate(processes_for(unit), 2):
    for j, k in enumerate(['Unit', 'Process', 'Machines', 'Capacity_per_day', 'Active'], 1):
        c = mp.cell(i, j, p[k]); c.font = F(c='0000FF' if j >= 3 else '222222')
        if j >= 3: c.fill = fill(INP)
dv2 = DataValidation(type='list', formula1='"Y,N"'); mp.add_data_validation(dv2); dv2.add(f'E2:E{max(mp.max_row, 30)}')

# ---- Lists
ls = wb.create_sheet('Lists'); head(ls, ['Hour slots', 'Wash types', 'Wash downtime reasons', 'Dry downtime reasons', 'Styles (add active styles)', 'Dry processes'],
                                   widths={'Styles (add active styles)': 30, 'Wash downtime reasons': 26, 'Dry downtime reasons': 26, 'Dry processes': 28})
for i, v in enumerate(hour_labels(8), 2): ls.cell(i, 1, v)
for col, items in [(2, WASH_TYPES), (3, WASH_REASONS), (4, DRY_REASONS), (6, [p['Process'] for p in processes_for(unit)])]:
    for i, v in enumerate(items, 2): ls.cell(i, col, v)
for i, v in enumerate(['ST-4412 Dark Stone', 'ST-5120 Acid Blast'], 2):
    c = ls.cell(i, 5, v); c.fill = fill(INP); c.font = F(c='0000FF')

# ---- Logs
LAST = 60000
def validations(ws, specs):
    for formula, rng, strict in specs:
        d = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=strict); ws.add_data_validation(d); d.add(rng)
    d = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True, errorTitle='Check value', error='Enter a number >= 0')
    return d

wl = wb.create_sheet('Wash_Log')
head(wl, ['Date', 'Hour', 'Unit', 'Machine_ID', 'Wash_Type', 'Style', 'Input', 'Output', 'QC_Inspected', 'Defect', 'Rewash', 'Downtime_Min', 'Downtime_Reason', 'Remarks'],
     key_n=4, widths={'Hour': 13, 'Machine_ID': 16, 'Style': 26, 'Downtime_Reason': 24, 'Remarks': 18})
ex = ['2026-10-08', '08:00-09:00', unit, machines_for(unit)[0]['Machine_ID'], 'Denim', 'ST-4412 Dark Stone', 120, 110, 90, 2, 1, '', '', 'EXAMPLE - delete me']
for j, v in enumerate(ex, 1):
    c = wl.cell(2, j, v); c.font = F(c='888888' if j <= 4 else '0000FF')
    if j <= 4: c.fill = fill(KEY)
validations(wl, [(f'=Lists!$B$2:$B$10', f'E2:E{LAST}', True), (f'=Lists!$E$2:$E$400', f'F2:F{LAST}', False), (f'=Lists!$C$2:$C$12', f'M2:M{LAST}', True)])
dn = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True, errorTitle='Check value', error='Enter a number >= 0'); wl.add_data_validation(dn); dn.add(f'G2:L{LAST}')

dl = wb.create_sheet('Dry_Log')
head(dl, ['Date', 'Hour', 'Unit', 'Process', 'Style', 'Input', 'Output', 'Defect', 'Rework', 'Downtime_Min', 'Downtime_Reason', 'Remarks'],
     key_n=4, widths={'Hour': 13, 'Process': 28, 'Style': 26, 'Downtime_Reason': 24, 'Remarks': 18})
ex = ['2026-10-08', '08:00-09:00', unit, processes_for(unit)[2]['Process'] if processes_for(unit) else 'Pattern Whisker', 'ST-4412 Dark Stone', 900, 880, 15, 8, '', '', 'EXAMPLE - delete me']
for j, v in enumerate(ex, 1):
    c = dl.cell(2, j, v); c.font = F(c='888888' if j <= 4 else '0000FF')
    if j <= 4: c.fill = fill(KEY)
validations(dl, [(f'=Lists!$E$2:$E$400', f'E2:E{LAST}', False), (f'=Lists!$D$2:$D$10', f'K2:K{LAST}', True)])
dn2 = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True, errorTitle='Check value', error='Enter a number >= 0'); dl.add_data_validation(dn2); dn2.add(f'F2:J{LAST}')

wp = wb.create_sheet('WIP_Log'); head(wp, ['Date', 'Unit', 'Stage', 'Opening_WIP', 'Remarks'], widths={'Stage': 28, 'Remarks': 30})
for j, v in enumerate(['2026-10-08', unit, 'WASH', 100000, 'EXAMPLE - delete me'], 1): wp.cell(2, j, v).font = F(c='0000FF')
d3 = DataValidation(type='list', formula1='=Lists!$F$2:$F$30', allow_blank=True, showErrorMessage=False); wp.add_data_validation(d3); d3.add('C3:C5000')
for ws in (wl, dl, wp): ws.column_dimensions['A'].width = 12
wb.save(out); print('saved', out)

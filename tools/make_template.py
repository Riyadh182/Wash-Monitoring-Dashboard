"""Builds the Google-Sheet template (.xlsx) for one unit: python make_template.py CWL-1 ../Wash_Pilot_Template_CWL-1.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL
from common import *

unit = sys.argv[1] if len(sys.argv) > 1 else 'CWL-1'
out = sys.argv[2] if len(sys.argv) > 2 else f'Wash_Pilot_Template_{unit}.xlsx'
if not CFG['units'][unit]['dry_verified'] and CFG['units'][unit]['dry_processes']:
    print(f'WARNING: dry-process capacity for {unit} duplicates an earlier unit in Capacity_August_2026 - verify before use.')
if not CFG['units'][unit]['dry_processes']:
    print(f'NOTE: no dry-process capacity in source for {unit}; Master_Process will be empty - fill it in.')

F = lambda b=False, c='222222', s=10: Font(name='Arial', size=s, bold=b, color=c)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
NAVY, TEAL, KEY, INP = '1B2A41', '2A9D8F', 'ECEFF3', 'FFF8CC'
wb = Workbook()

def head(ws, cols, key_n=0, widths=None, row=1):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row, j, h); c.font = F(True, 'FFFFFF', 10); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.fill = fill(NAVY if j <= key_n else TEAL)
        ws.column_dimensions[CL(j)].width = (widths or {}).get(h, 16)
    ws.row_dimensions[row].height = 32; ws.freeze_panes = ws.cell(row + 1, 1)

# ---- README
r = wb.active; r.title = 'README'; r.sheet_view.showGridLines = False
r.column_dimensions['A'].width = 3; r.column_dimensions['B'].width = 120
lines = [('WASH & DRY HOURLY LOG - ' + unit, True), ('', False),
 ('1. Upload this file to Google Drive and open with Google Sheets (File > Save as Google Sheets).', False),
 ('2. Extensions > Apps Script > paste apps_script/Code.gs > Save > reload. A "Wash Tools" menu appears.', False),
 ('3. Each production day: Wash Tools > Create rows for a day. It fills Date, Hour, Unit and Machine/Process for all 24 hours.', False),
 ('4. Supervisors type only numbers (Input, Output, QC, Defect, Rewash, Downtime). Grey columns are keys - do not edit.', False),
 ('5. Style, PO and Color: pick from the dropdown (comes from Order_Master) or type. Wash type and Downtime reason are dropdowns.', False),
 ('5b. Order_Master: one row per Style + PO + Color (buyer, CRD, order qty, shade dates, merchant). Plan_Start..Plan_End = wash INPUT window; Order_Qty is spread evenly over it to get plan till today. Feeds the dashboard Style details table. Write ALL dates as yyyy-mm-dd.', False),
 ('5c. Daily_Plan (optional): real planned input and output pcs per date per order. Without it the dashboard spreads Order_Qty evenly over Plan_Start..Plan_End.', False),
 ('6. WIP_Log: one row per stage per day with the physical opening WIP. Stage = WASH or a dry process name (exactly as in Master_Process).', False),
 ('7. Master_Machine / Master_Process: plan per hour comes from here. Edit Plan_per_hr for a machine, or set Active = N for a machine that is down for the day.', False),
 ('8. Share this Google Sheet with the dashboard service account as VIEWER. Floor staff get EDITOR access.', False),
 ('9. Month end: File > Make a copy (name it with the month), keep it as the archive, then clear Wash_Log/Dry_Log/WIP_Log data rows (keep headers) for the new month.', False),
 ('', False), ('Rows with "EXAMPLE" in Remarks are ignored by the dashboard - delete them whenever you like.', False),
 ('Hour slots assume the production day starts at Day_Start_Hour in Settings (08:00 default): hours after midnight belong to the same production date.', False)]
for i, (t, b) in enumerate(lines, 2):
    c = r.cell(i, 2, t); c.font = F(b, NAVY if b else '222222', 14 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')

# ---- Settings
s = wb.create_sheet('Settings'); head(s, ['Key', 'Value', 'Meaning'], widths={'Key': 20, 'Value': 12, 'Meaning': 90})
notes = {'Run_Hours': 'Hours the plant runs per day (hourly plan = capacity per day / this)', 'Day_Start_Hour': 'Production day start hour (0-23); shifts A/B/C are 8 hours each',
         'Plan_Basis': '1 = plan at 100% of capacity; 0.85 = plan at 85% load', 'Target_Pct': 'Achievement at/above = green', 'Alert_Pct': 'Achievement below = red',
         'Max_Defect_Pct': 'Defect limit (defect pcs / QC inspected pcs) - placeholder, set yours', 'Max_Rewash_Pct': 'Rewash limit (rewash pcs / output pcs) - placeholder, set yours', 'Output_Lag_Days': 'Style details: wash OUTPUT plan starts this many days after the input plan'}
for i, (k, v) in enumerate(DEFAULT_SETTINGS, 2):
    s.cell(i, 1, k).font = F(True); c = s.cell(i, 2, v); c.font = F(c='0000FF'); c.fill = fill(INP); s.cell(i, 3, notes[k]).font = F(s=9, c='666666')
    if k.endswith('Pct') or k == 'Plan_Basis': c.number_format = '0%'

# ---- Masters
mm = wb.create_sheet('Master_Machine'); head(mm, ['Unit', 'Machine_ID', 'Type', 'Plan_per_hr', 'Active'], key_n=3)
for i, m in enumerate(machines_for(unit), 2):
    for j, k in enumerate(['Unit', 'Machine_ID', 'Type', 'Plan_per_hr', 'Active'], 1):
        c = mm.cell(i, j, m[k]); c.font = F(c='0000FF' if j >= 4 else '222222')
        if j >= 4: c.fill = fill(INP)
dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=False); mm.add_data_validation(dv); dv.add(f'E2:E{mm.max_row}')
mm['G1'] = 'Plan_per_hr = unit wash capacity per day / machines / run hours (Capacity_August_2026). It is an average - overwrite per machine when you know the real capacity.'
mm['G1'].font = F(s=9, c='666666'); mm.column_dimensions['G'].width = 80

mp = wb.create_sheet('Master_Process'); head(mp, ['Unit', 'Process', 'Machines', 'Capacity_per_day', 'Active'], key_n=2, widths={'Process': 28})
for i, p in enumerate(processes_for(unit), 2):
    for j, k in enumerate(['Unit', 'Process', 'Machines', 'Capacity_per_day', 'Active'], 1):
        c = mp.cell(i, j, p[k]); c.font = F(c='0000FF' if j >= 3 else '222222')
        if j >= 3: c.fill = fill(INP)
dv2 = DataValidation(type='list', formula1='"Y,N"'); mp.add_data_validation(dv2); dv2.add(f'E2:E{max(mp.max_row, 30)}')

# ---- Lists
ls = wb.create_sheet('Lists'); head(ls, ['Hour slots', 'Wash types', 'Wash downtime reasons', 'Dry downtime reasons', 'Dry processes'],
                                   widths={'Wash downtime reasons': 26, 'Dry downtime reasons': 26, 'Dry processes': 28})
for i, v in enumerate(hour_labels(8), 2): ls.cell(i, 1, v)
for col, items in [(2, WASH_TYPES), (3, WASH_REASONS), (4, DRY_REASONS), (5, [p['Process'] for p in processes_for(unit)])]:
    for i, v in enumerate(items, 2): ls.cell(i, col, v)

# ---- Order master
om = wb.create_sheet('Order_Master')
OC = ['Wash_Plant', 'Buyer', 'Style', 'PO', 'Color', 'Wash_Type', 'CRD', 'CRD_Revisions', 'Updated_CRD', 'Sewing_Unit', 'Order_Qty', 'Plan_Start', 'Plan_End',
      'Shade_Submit_Date', 'Shade_Approval_Date', 'Merchant', 'Remarks']
head(om, OC, key_n=5, widths={'Style': 26, 'Merchant': 22, 'Shade_Submit_Date': 16, 'Shade_Approval_Date': 18, 'Remarks': 22})
import datetime
exo = [unit, 'Buyer name', 'ST-4412 Dark Stone', 'PO-0001', 'Dark Blue', 'Denim', '2026-11-20', 0, '', 'TISWL-01', 20000, '2026-10-08', '2026-10-20', '2026-10-05', '', 'Merchant name', 'EXAMPLE - delete me']
for j, v in enumerate(exo, 1):
    c = om.cell(2, j, v); c.font = F(c='0000FF')
for r_ in range(2, 2001):
    for j in (7, 9, 12, 13, 14, 15): om.cell(r_, j).number_format = 'yyyy-mm-dd'
dvw = DataValidation(type='list', formula1='=Lists!$B$2:$B$10', allow_blank=True, showErrorMessage=False); om.add_data_validation(dvw); dvw.add('F2:F2000')

# ---- Logs
LAST = 60000
def listval(ws, formula, rng, strict):
    d = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=strict); ws.add_data_validation(d); d.add(rng)
def numval(ws, rng):
    d = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True, errorTitle='Check value', error='Enter a number >= 0')
    ws.add_data_validation(d); d.add(rng)

wl = wb.create_sheet('Wash_Log')
head(wl, ['Date', 'Hour', 'Unit', 'Machine_ID', 'Wash_Type', 'Style', 'PO', 'Color', 'Input', 'Output', 'QC_Inspected', 'Defect', 'Rewash', 'Rewash_Output', 'Reject',
          'Downtime_Min', 'Downtime_Reason', 'Remarks'],
     key_n=4, widths={'Hour': 13, 'Machine_ID': 16, 'Style': 26, 'PO': 12, 'Color': 16, 'Downtime_Reason': 24, 'Remarks': 18})
ex = ['2026-10-08', '08:00-09:00', unit, machines_for(unit)[0]['Machine_ID'], 'Denim', 'ST-4412 Dark Stone', 'PO-0001', 'Dark Blue', 120, 110, 90, 2, 1, 0, 0, '', '', 'EXAMPLE - delete me']
for j, v in enumerate(ex, 1):
    c = wl.cell(2, j, v); c.font = F(c='888888' if j <= 4 else '0000FF')
    if j <= 4: c.fill = fill(KEY)
listval(wl, '=Lists!$B$2:$B$10', f'E2:E{LAST}', True)
listval(wl, '=Order_Master!$C$2:$C$2000', f'F2:F{LAST}', False)
listval(wl, '=Order_Master!$D$2:$D$2000', f'G2:G{LAST}', False)
listval(wl, '=Order_Master!$E$2:$E$2000', f'H2:H{LAST}', False)
listval(wl, '=Lists!$C$2:$C$12', f'Q2:Q{LAST}', True)
numval(wl, f'I2:P{LAST}')

dl = wb.create_sheet('Dry_Log')
head(dl, ['Date', 'Hour', 'Unit', 'Process', 'Style', 'PO', 'Color', 'Input', 'Output', 'Defect', 'Rework', 'Downtime_Min', 'Downtime_Reason', 'Remarks'],
     key_n=4, widths={'Hour': 13, 'Process': 28, 'Style': 26, 'PO': 12, 'Color': 16, 'Downtime_Reason': 24, 'Remarks': 18})
ex = ['2026-10-08', '08:00-09:00', unit, processes_for(unit)[2]['Process'] if processes_for(unit) else 'Pattern Whisker', 'ST-4412 Dark Stone', 'PO-0001', 'Dark Blue', 900, 880, 15, 8, '', '', 'EXAMPLE - delete me']
for j, v in enumerate(ex, 1):
    c = dl.cell(2, j, v); c.font = F(c='888888' if j <= 4 else '0000FF')
    if j <= 4: c.fill = fill(KEY)
listval(dl, '=Order_Master!$C$2:$C$2000', f'E2:E{LAST}', False)
listval(dl, '=Order_Master!$D$2:$D$2000', f'F2:F{LAST}', False)
listval(dl, '=Order_Master!$E$2:$E$2000', f'G2:G{LAST}', False)
listval(dl, '=Lists!$D$2:$D$10', f'M2:M{LAST}', True)
numval(dl, f'H2:L{LAST}')

dpn = wb.create_sheet('Daily_Plan')
head(dpn, ['Date', 'Wash_Plant', 'Style', 'PO', 'Color', 'Plan_Input', 'Plan_Output'], key_n=5, widths={'Style': 26, 'PO': 12, 'Color': 16})
dpn['I1'] = 'Optional. One row per date per order: how many pcs you PLAN to input and to output that day. Orders listed here override the straight-line plan from Order_Master. Dates as yyyy-mm-dd.'
dpn['I1'].font = F(s=9, c='666666'); dpn.column_dimensions['I'].width = 90
for r_ in range(2, 5001): dpn.cell(r_, 1).number_format = 'yyyy-mm-dd'

wp = wb.create_sheet('WIP_Log'); head(wp, ['Date', 'Unit', 'Stage', 'Opening_WIP', 'Remarks'], widths={'Stage': 28, 'Remarks': 30})
for j, v in enumerate(['2026-10-08', unit, 'WASH', 100000, 'EXAMPLE - delete me'], 1): wp.cell(2, j, v).font = F(c='0000FF')
d3 = DataValidation(type='list', formula1='=Lists!$E$2:$E$30', allow_blank=True, showErrorMessage=False); wp.add_data_validation(d3); d3.add('C3:C5000')
for ws in (wl, dl, wp): ws.column_dimensions['A'].width = 12
wb.save(out); print('saved', out)

"""Download helpers (CSV and formatted Excel)."""
from __future__ import annotations

import io

import pandas as pd


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")      # utf-8-sig opens correctly in Excel


def to_excel_bytes(sheets: dict) -> bytes:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    thin = Side(style="thin", color="FFCBD5E1")
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            nm = str(name)[:31].replace("/", "-").replace("\\", "-").replace("[", "(").replace("]", ")").replace(":", "-").replace("*", "").replace("?", "")
            df.to_excel(xw, sheet_name=nm, index=False)
            ws = xw.sheets[nm]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFFFF")
                c.fill = PatternFill("solid", fgColor="FF4F46E5")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 32
            for i, col in enumerate(df.columns, start=1):
                width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200).tolist())) if len(df) else len(str(col))
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 45)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    if isinstance(c.value, float):
                        c.number_format = "#,##0.0" if abs(c.value) < 1000 else "#,##0"
            ws.freeze_panes = "A2"
    return buf.getvalue()

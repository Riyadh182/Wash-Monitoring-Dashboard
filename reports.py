"""Report builders and Excel export (pure pandas / openpyxl - no Streamlit)."""
import io
import numpy as np
import pandas as pd
import calc

LABELS = {
    'Wash_Plant': 'Wash plant', 'Sewing_Unit': 'Sewing unit', 'Buyer': 'Buyer', 'Style': 'Style', 'PO': 'PO', 'Color': 'Color', 'Wash_Type': 'Wash type',
    'Merchant': 'Merchant', 'Eff_CRD': 'CRD (effective)', 'CRD': 'CRD', 'Updated_CRD': 'Updated CRD', 'Days_to_CRD': 'Days to CRD', 'Next_CRD': 'Next CRD (open)',
    'Order_Qty': 'Order qty', 'Orders': 'Orders', 'Styles': 'Styles', 'Input_Plan': 'Plan input till today', 'Input': 'Actual input (received)',
    'Input_Var': 'Input variation', 'Input_Achv_Pct': 'Input achievement (%)', 'Output_Plan': 'Plan output till today', 'Output': 'Actual output (delivered)',
    'Output_Var': 'Output variation', 'Achv_Pct': 'Achievement (%)', 'WIP': 'WIP', 'Rewash': 'Rewash input', 'Rewash_Output': 'Rewash output', 'Rewash_WIP': 'Rewash WIP',
    'Total_WIP': 'Total WIP', 'Rewash_Pct': 'Rewash (%)', 'Defect': 'Quality issue (pcs)', 'QC_Inspected': 'QC inspected', 'DHU': 'DHU (%)', 'Reject': 'Rejection (pcs)',
    'Reject_Pct': 'Rejection (%)', 'Balance': 'Balance to deliver', 'Pending_Input': 'Not yet received', 'Crisis_Orders': 'Crisis orders (CRD <= 7d)', 'Crisis_Pcs': 'Crisis pcs',
    'Required_Rate': 'Required pcs/day', 'Rate3': 'Last-3-day output/day', 'Risk': 'Risk', 'Where_Stuck': 'Where stuck', 'Can_Finish': 'Can finish by CRD?',
    'CRD_Group': 'CRD', 'Status': 'Status', 'Date': 'Date', 'Opening_WIP': 'Opening WIP', 'Closing_WIP': 'Closing WIP', 'Change': 'Change vs opening',
    'WIP_Yesterday': 'WIP yesterday', 'WIP_Today': 'WIP today', 'Receive_Today': 'Receive today (input)', 'Delivery_Today': 'Delivery today (output)',
    'Change_Pct': 'Change (%)', 'Machines_Full': 'Machine IDs', 'Prior_Input': 'Input before logging', 'Prior_Output': 'Output before logging',
    'Month': 'Month', 'Machine_ID': 'Machine', 'Process': 'Process', 'Type': 'Type', 'Cap_hr': 'Capacity / hr', 'Days': 'Days', 'Avail_Hours': 'Available hours',
    'Capacity': 'Capacity (pcs)', 'Allotted': 'Allotted plan (pcs)', 'Idle_Capacity': 'Idle capacity (pcs)', 'Allot_Pct': 'Allotment (%)',
    'Plan_Achv_Pct': 'Plan achievement (%)', 'Util_Pct': 'Capacity utilisation (%)', 'Downtime_Min': 'Downtime (min)', 'Down_Pct': 'Downtime (% of time)',
    'Lost_Pcs': 'Lost pcs (downtime)', 'Defect_Pct': 'Defect (%)', 'Rework_Pct': 'Rework (%)', 'Top_Downtime_Reasons': 'Top downtime reasons',
    'Overrun_Orders': 'Plan-overrun orders', 'Plan_End': 'Wash plan end', 'Overrun_Days': 'Days past plan end', 'Flags': 'Red flags', 'Priority': 'Wash priority (EDD, FIFO)',
    'Forecast_Date': 'Forecast completion', 'Late_Days': 'Late (days) vs CRD', 'Verdict': 'Forecast verdict', 'Plan_Revisions': 'Plan revisions', 'Edited': 'Edited entries',
    'First_Receive_Date': 'First received', 'Sew_Delay_Days': 'Sewing delay (days)', 'Suggested_Wash': 'Suggested wash (pcs)', 'Item': 'Item', 'Season': 'Season', 'Missing_Body': 'Missing body',
    'Wet_Capacity_Day': 'Wet (wash) capacity / day', 'Dry_Capacity_Day': 'Dry capacity / day',
}
ORDER_VIEW = ['Wash_Plant', 'Sewing_Unit', 'Buyer', 'Style', 'PO', 'Color', 'Wash_Type', 'Eff_CRD', 'Days_to_CRD', 'Order_Qty', 'Input_Plan', 'Input', 'Input_Var',
              'Output_Plan', 'Output', 'Output_Var', 'Achv_Pct', 'WIP', 'Rewash_WIP', 'Total_WIP', 'Balance']
GROUP_VIEW = ['Orders', 'Order_Qty', 'Input_Plan', 'Input', 'Input_Var', 'Input_Achv_Pct', 'Output_Plan', 'Output', 'Output_Var', 'Achv_Pct', 'WIP', 'Rewash_WIP',
              'Total_WIP', 'Balance', 'Rewash', 'Rewash_Pct', 'DHU', 'Reject', 'Reject_Pct', 'Crisis_Orders', 'Crisis_Pcs', 'Overrun_Orders', 'Next_CRD']
CRISIS_VIEW = ['Risk', 'Wash_Plant', 'Sewing_Unit', 'Buyer', 'Style', 'PO', 'Color', 'Eff_CRD', 'Days_to_CRD', 'Order_Qty', 'Output', 'Balance', 'Total_WIP', 'Pending_Input',
               'Rate3', 'Required_Rate', 'Can_Finish', 'Where_Stuck', 'Plan_End', 'Overrun_Days', 'Flags', 'Merchant']


def cols(df, want):
    return df[[c for c in want if c in df.columns]]


def export_frame(df):
    d = df.copy()
    for c in d.columns:
        if pd.api.types.is_datetime64_any_dtype(d[c]):
            d[c] = d[c].dt.strftime('%Y-%m-%d')
        elif pd.api.types.is_float_dtype(d[c]):
            d[c] = d[c].round(1)
        elif pd.api.types.is_bool_dtype(d[c]):
            d[c] = np.where(d[c], 'Yes', 'No')
    d = d.rename(columns=LABELS)
    if d.columns.duplicated().any():     # keep names unique so Excel / pandas do not choke
        seen, cols_ = {}, []
        for c in d.columns:
            seen[c] = seen.get(c, 0) + 1
            cols_.append(c if seen[c] == 1 else f'{c} ({seen[c]})')
        d.columns = cols_
    return d


def excel_bytes(sheets):
    """sheets: {name: DataFrame}. Styled header, frozen header row, auto widths, thousands separators."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xw:
        used = set()
        for name, df in sheets.items():
            nm = ''.join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or 'Sheet'
            while nm in used:
                nm = nm[:28] + '_' + str(len(used))
            used.add(nm)
            e = export_frame(df) if df is not None and len(df.columns) else pd.DataFrame({'info': ['No data']})
            e.to_excel(xw, sheet_name=nm, index=False)
            ws = xw.sheets[nm]
            fill = PatternFill('solid', start_color='1D4ED8', end_color='1D4ED8')
            for j, c in enumerate(e.columns, 1):
                h = ws.cell(1, j); h.font = Font(bold=True, color='FFFFFF'); h.fill = fill
                h.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
                q90 = e[c].astype(str).str.len().quantile(.9) if len(e) else 8
                width = max(len(str(c)) * 0.9, min(40, (int(q90) if pd.notna(q90) else 8) + 2), 9)
                ws.column_dimensions[get_column_letter(j)].width = min(width, 42)
                if pd.api.types.is_numeric_dtype(e[c]):
                    for r in range(2, len(e) + 2):
                        ws.cell(r, j).number_format = '#,##0.0' if pd.api.types.is_float_dtype(e[c]) and e[c].abs().max() < 1000 else '#,##0'
            ws.row_dimensions[1].height = 32
            ws.freeze_panes = 'A2'
    return buf.getvalue()


def buyer_crisis(C):
    if C is None or C.empty:
        return pd.DataFrame()
    g = C.groupby('Buyer').agg(Orders=('Style', 'size'), Overdue=('Risk', lambda s: int((s == 'Overdue').sum())), Critical=('Risk', lambda s: int((s == 'Critical').sum())),
                               Next_CRD=('Eff_CRD', 'min'), Order_Qty=('Order_Qty', 'sum'), Balance=('Balance', 'sum'), Total_WIP=('Total_WIP', 'sum'),
                               Pending_Input=('Pending_Input', 'sum'), Required_Rate=('Required_Rate', 'sum'), Rate3=('Rate3', 'sum')).reset_index()
    return g.sort_values('Balance', ascending=False).reset_index(drop=True)


def report_set(S, pdaily, asof, horizon=7, d0=None, d1=None):
    """All regular reports as {sheet name: DataFrame} for the (already filtered) order status S and plant daily table."""
    R = {}
    asof = pd.Timestamp(asof).normalize()
    if pdaily is not None and len(pdaily):
        R['1 WIP vs yesterday'] = calc.wip_compare(pdaily, asof)
        pd_ = pdaily if d0 is None else pdaily[(pdaily['Date'] >= pd.Timestamp(d0)) & (pdaily['Date'] <= pd.Timestamp(d1))]
        R['2 Daily receive-delivery'] = pd_
        tot = pd_.groupby('Date')[['Input', 'Output', 'Rewash', 'Rewash_Output', 'Opening_WIP', 'Closing_WIP']].sum().reset_index()
        R['2b Daily total all plants'] = tot
    if S is not None and len(S):
        R['3 Plant summary'] = cols(calc.group_summary(S, ['Wash_Plant'], horizon), ['Wash_Plant'] + GROUP_VIEW)
        R['4 Sewing unit summary'] = cols(calc.group_summary(S, ['Sewing_Unit', 'Wash_Plant'], horizon), ['Sewing_Unit', 'Wash_Plant'] + GROUP_VIEW)
        R['5 Sewing unit style detail'] = cols(S.sort_values(['Sewing_Unit', 'Wash_Plant', 'Buyer', 'Style']), ORDER_VIEW)
        R['6 Buyer summary'] = cols(calc.group_summary(S, ['Buyer'], horizon), ['Buyer'] + GROUP_VIEW)
        R['7 Buyer style detail'] = cols(S.sort_values(['Buyer', 'Wash_Plant', 'Style']), ORDER_VIEW)
        cs = calc.crd_summary(S, asof, 'date')
        R['8 CRD wise'] = cols(cs, ['CRD_Group', 'Days_to_CRD', 'Status'] + GROUP_VIEW)
        C = calc.crisis_table(S, horizon)
        R['9 CRD crisis'] = cols(C, CRISIS_VIEW) if len(C) else pd.DataFrame()
        R['10 Buyer crisis'] = buyer_crisis(C)
    return R
